"""Join public run identities to the published workbook without guessing pairing."""

import argparse
import csv
import json
import re
from pathlib import Path
from openpyxl import load_workbook


def build(workbook, metadata, output):
    book = load_workbook(workbook, read_only=True, data_only=True)
    published = {}
    for r in book.worksheets[0].iter_rows(min_row=2, max_col=11, values_only=True):
        if r[0] and isinstance(r[0], str) and ".out.counts" in r[0]:
            key = str(r[1])
            if key in published:
                raise ValueError(f"Duplicate published vial {key}")
            published[key] = r
    groups = {}
    group = None
    for (value,) in book["DEseq group annotation"].iter_rows(
        max_col=1, values_only=True
    ):
        if not value:
            continue
        if value.startswith("Group "):
            group = value
        else:
            match = re.search(r"_(JA\d+|5553)_", value)
            if not match or group is None:
                raise ValueError(value)
            groups.setdefault(match[1], []).append(group)
    output.mkdir(parents=True, exist_ok=True)
    rows = []
    for run in json.loads(metadata.read_text()):
        attrs = run["attributes"]
        vial = attrs["Vial Number"]
        p = published[vial]
        treatment = {"PaceMakerCon": "Control", "PaceMakerTre": "Treatment"}[
            attrs["Experimental Group"]
        ]
        assert p[2] == treatment and int(str(attrs["Age In Days"]).rstrip("d")) == p[6]
        flags = []
        if attrs["Stage"] != p[5]:
            flags.append("stage_labels_differ_in_specificity_or_value")
        if vial not in groups:
            flags.append("not_in_published_comparison_groups")
        rows.append(
            dict(
                run=run["run_accession"],
                study=run["study_accession"],
                biosample=run["sample_accession"],
                vial=vial,
                genotype=attrs["Colony Id"],
                treatment=treatment,
                hours_after_treatment=p[3],
                age_days=p[6],
                archive_stage=attrs["Stage"],
                published_stage=p[5],
                published_groups=";".join(groups.get(vial, [])),
                archive_paired_records=run["read_count"],
                archive_bytes=sum(map(int, run["fastq_bytes"].split(";"))),
                proposed_pair="unresolved",
                inference_status="blocked_pending_experimental_unit_review",
                flags=";".join(flags) if flags else "none",
            )
        )
    assert len({r["run"] for r in rows}) == len(rows) == len(published)
    with (output / "samples.tsv").open("w") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]), delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)
    summary = dict(
        runs=len(rows),
        biosamples=len({r["biosample"] for r in rows}),
        published_comparison_samples=sum(bool(r["published_groups"]) for r in rows),
        archive_bytes=sum(r["archive_bytes"] for r in rows),
        published_times_hours=sorted({r["hours_after_treatment"] for r in rows}),
        stage_label_differences=sum("stage_labels_differ" in r["flags"] for r in rows),
        genotype_counts={
            g: sum(r["genotype"] == g for r in rows)
            for g in sorted({r["genotype"] for r in rows})
        },
    )
    (output / "sample-audit.json").write_text(json.dumps(summary, indent=2) + "\n")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("workbook", type=Path)
    p.add_argument("metadata", type=Path)
    p.add_argument("output", type=Path)
    a = p.parse_args()
    build(a.workbook, a.metadata, a.output)
