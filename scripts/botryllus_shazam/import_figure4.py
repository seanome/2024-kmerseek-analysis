"""Extract the published Figure 4 inputs without inferring missing results."""

import argparse, csv, gzip, hashlib, json, re
from collections import Counter
from pathlib import Path
import openpyxl

p = argparse.ArgumentParser()
p.add_argument("workbook", type=Path)
p.add_argument("samples", type=Path)
p.add_argument("output", type=Path)
a = p.parse_args()
a.output.mkdir(parents=True, exist_ok=True)
w = openpyxl.load_workbook(a.workbook, read_only=True, data_only=True)
meta = list(csv.DictReader(a.samples.open(), delimiter="\t"))
by_vial = {r["vial"]: r for r in meta}
assert len(by_vial) == 64


def write(name, columns, rows):
    path = a.output / name
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "wt") as f:
        t = csv.writer(f, delimiter="\t")
        t.writerow(columns)
        t.writerows(rows)


groups = []
group = None
for (value,) in w["DEseq group annotation"].iter_rows(values_only=True):
    if not value:
        continue
    if value.startswith("Group "):
        group = value
        continue
    vial = value.split("_")[1]
    assert vial in by_vial
    groups.append(dict(group=group, run=by_vial[vial]["run"], vial=vial, sample=value))
write(
    "published-groups.tsv",
    ["group", "run", "vial", "sample"],
    [[r[k] for k in ["group", "run", "vial", "sample"]] for r in groups],
)
sets = []
for row in w["Fig 4_M1 M2 Botryllus gene set"].iter_rows(min_row=3, values_only=True):
    for col, role in [(0, "proimmune"), (2, "cell_repair")]:
        if row[col]:
            gid, symbol = row[col].split(" ", 1)
            assert re.fullmatch("g[0-9]+", gid)
            sets.append([gid, symbol, role])
assert Counter(x[2] for x in sets) == dict(proimmune=13, cell_repair=133)
assert len({x[0] for x in sets}) == 146
write("immune-gene-sets.tsv", ["gene_id", "gene_name", "role"], sets)
counts = w["Pacemaker_Counts_Table"]
it = counts.iter_rows(values_only=True)
header = next(it)
runs = []
for name in header[3:]:
    vial = name.split("_")[1]
    runs.append(by_vial[vial]["run"])
assert len(runs) == len(set(runs)) == 64
labels = []
totals = [0] * 64
seen = set()
with gzip.open(a.output / "published-counts.tsv.gz", "wt") as f:
    t = csv.writer(f, delimiter="\t")
    t.writerow(["tag"] + runs)
    for row in it:
        if not row[0]:
            continue
        gid = str(row[0])
        assert gid not in seen
        seen.add(gid)
        if gid.startswith("__") or gid == "Sum":
            continue
        assert re.fullmatch("g[0-9]+", gid), gid
        vals = list(row[3:])
        assert len(vals) == 64 and all(
            isinstance(x, (int, float)) and x >= 0 and x == int(x) for x in vals
        )
        totals = [x + y for x, y in zip(totals, vals)]
        t.writerow([gid] + [int(x) for x in vals])
        labels.append(row[:3])
write("gene-labels.tsv", ["gene_id", "gene_name", "description"], labels)
assert {x[0] for x in sets} <= set(x[0] for x in labels)
write(
    "library-count-totals.tsv",
    ["run", "assigned_gene_counts", "passes_10000"],
    [[r, int(n), n >= 10000] for r, n in zip(runs, totals)],
)
stats = []
for sheet in w.sheetnames[3:9]:
    rows = w[sheet].iter_rows(values_only=True)
    hdr = list(next(rows)[:11])
    assert hdr[0] == "Gene_ID" and hdr[9:11] == ["pvalue", "padj"]
    n = raw = adj = 0
    finite = 0
    with gzip.open(a.output / (sheet + ".tsv.gz"), "wt") as f:
        t = csv.writer(f, delimiter="\t")
        t.writerow(hdr)
        for row in rows:
            if not row[0]:
                continue
            row = row[:11]
            n += 1
            t.writerow(row)
            if isinstance(row[9], (float, int)) and isinstance(row[6], (float, int)):
                finite += 1
                raw += row[9] <= 0.05 and abs(row[6]) > 1
            if isinstance(row[10], (float, int)) and isinstance(row[6], (float, int)):
                adj += row[10] <= 0.05 and abs(row[6]) > 1
    stats.append(
        dict(
            sheet=sheet,
            genes=n,
            finite_p=finite,
            raw_p_005_abs_lfc_gt1=raw,
            adjusted_p_005_abs_lfc_gt1=adj,
        )
    )
# Preserve every declared hub count and the original gene-list text separately.
cohort = None
times = []
bubbles = []
for row in w["fig4a&bgene_count"].iter_rows(values_only=True):
    if row[0] and str(row[0]).startswith("Fig 4"):
        cohort = "young" if "Young" in row[0] else "old"
        continue
    if row[0] == "Functional Regulation":
        times = [x for x in row[2:] if x]
        continue
    if row[1] in ("Up-reg", "Down-reg"):
        for time, n in zip(times, row[2:]):
            assert isinstance(n, int)
            bubbles.append(
                dict(cohort=cohort, hub=row[0], direction=row[1], time=time, count=n)
            )
write(
    "published-hub-counts.tsv",
    ["cohort", "hub", "direction", "time", "count"],
    [[r[k] for k in ["cohort", "hub", "direction", "time", "count"]] for r in bubbles],
)
hubs = []
for sheet, cohort in [
    ("Temporal Reg gene list young ", "young"),
    ("Temporal Reg gene list Old", "old"),
]:
    rows = w[sheet].iter_rows(values_only=True)
    hdr = next(rows)
    for row in rows:
        if row[1] not in ("Up-reg", "Down-reg"):
            continue
        for i in range(3, 10, 2):
            if i >= len(row) or not row[i]:
                continue
            cell = str(row[i])
            m = re.match(r"count\s*=?\s*(\d+)\s*[:;]?\s*(.*)", cell, re.I)
            if not m:
                hubs.append(
                    dict(
                        cohort=cohort,
                        hub=row[0],
                        direction=row[1],
                        time=hdr[i],
                        declared_count=None,
                        gene_text=cell,
                        parse_status="not a numeric gene-count cell",
                    )
                )
                continue
            hubs.append(
                dict(
                    cohort=cohort,
                    hub=row[0],
                    direction=row[1],
                    time=hdr[i],
                    declared_count=int(m[1]),
                    gene_text=m[2],
                )
            )
(a.output / "published-hub-gene-text.json").write_text(
    json.dumps(hubs, indent=2) + "\n"
)
manifest = dict(
    source_url="https://www.ebi.ac.uk/europepmc/webservices/rest/PMC13229227/supplementaryFiles",
    workbook_sha256=hashlib.sha256(a.workbook.read_bytes()).hexdigest(),
    sheets=w.sheetnames,
    genes=len(labels),
    libraries=len(runs),
    gene_set_counts=dict(Counter(x[2] for x in sets)),
    published_groups=dict(Counter(r["group"] for r in groups)),
    low_count_libraries=[r for r, n in zip(runs, totals) if n < 10000],
    published_de_tables=stats,
    hub_count_cells=len(bubbles),
    status="Published inputs extracted; no raw-read or adjusted-model reproduction inferred",
)
(a.output / "source-manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
print(json.dumps(manifest, indent=2))
