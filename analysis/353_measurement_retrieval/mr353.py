#!/usr/bin/env python3
"""Notebook 353: does reduced-alphabet search help a scientist find existing measurements
on a short protein stretch faster than a conventional workflow?

One command-line tool with one subcommand per step. Every step reads and writes files in
--work (default: ./work), so each can be rerun alone. The rules (settings grid, pass rule,
coordinates, overlap) are in evaluation_plan.yaml; this file only carries them out.

Steps, in order:

  import-evidence   ELM instances + Sanborn et al. 2021 activation tiles -> evidence table,
                    one FASTA of the proteins that carry them, and a manifest of checksums
  import-queries    the analyst's queue (CSV, see templates/) -> development and evaluation
                    queries, split by arrival order before anything is searched
  make-controls     20 composition shuffles and 20 adjacent-pair shuffles of every query
  search-b          workflow B: exact full-alphabet substring lookup + local BLAST
  kmerseek-jobs     workflow A: writes the kmerseek index/search commands for Sherlock
  read-kmerseek     workflow A: reads kmerseek's CSV output into the common hit table
  cards             hits -> evidence cards (deduplicated, self-records removed), with the
                    shuffle-based expected count for every card
  tune              choose one alphabet, one k and every threshold on development queries
  freeze            write evaluation_plan.lock.json: checksums of every input and choice
  session-pack      the timed-session sheet: workflow order per evaluation query
  review-pack       the blind review sheet (after the timed sessions) and its sealed key
  score             per-query outcomes, the three pass criteria, PASS / FAIL /
                    INSUFFICIENT_EVIDENCE
  viewer            a single local HTML page showing every card

Coordinates are 0-based and end-exclusive everywhere inside this tool (kmerseek's
convention). Source coordinates are converted once, at import, and the original values are
kept in their own columns.

Run with the 2025-kmerseek-analysis environment (polars, openpyxl, pyyaml, numpy).
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import io
import json
import math
import os
import random
import re
import resource
import shutil
import statistics
import subprocess
import sys
import time
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

import polars as pl
import yaml

HERE = Path(__file__).resolve().parent
PLAN = HERE / "evaluation_plan.yaml"
AMINO = set("ACDEFGHIKLMNPQRSTVWY")

# ---- source files -------------------------------------------------------------------
ELM_INSTANCES_URL = "http://elm.eu.org/instances.tsv?q=*"
ELM_CLASSES_URL = "http://elm.eu.org/elms/elms_index.tsv"
ELM_FASTA_URL = "http://elm.eu.org/instances.fasta?q=*"
SANBORN_DOI = "10.7554/eLife.68068"
SANBORN_PMID = "33904398"
SANBORN_TILES_URL = (
    "https://cdn.elifesciences.org/articles/68068/elife-68068-fig1-data1-v3.xlsx"
)
UNIPROT_TSV = "https://rest.uniprot.org/uniprotkb/stream?format=tsv&fields=accession,sequence_version,sequence&query="
# --------------------------------------------------------------------------------------


def load_plan(path: Path = PLAN) -> dict:
    return yaml.safe_load(path.read_text())


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def md5(s: str) -> str:
    return hashlib.md5(s.encode()).hexdigest()


def read_fasta(path: Path) -> list[tuple[str, str]]:
    out, name, seq = [], None, []
    for line in Path(path).read_text().splitlines():
        if line.startswith(">"):
            if name is not None:
                out.append((name, "".join(seq)))
            name, seq = line[1:].strip(), []
        elif line.strip():
            seq.append(line.strip())
    if name is not None:
        out.append((name, "".join(seq)))
    return out


def write_fasta(records, path: Path) -> None:
    with open(path, "w") as f:
        for name, seq in records:
            f.write(f">{name}\n")
            for i in range(0, len(seq), 60):
                f.write(seq[i : i + 60] + "\n")


# ---- compute log --------------------------------------------------------------------
def log_compute(
    work: Path, step: str, t0: float, cpu0: float, extra: dict | None = None
):
    """Append wall time, CPU time (this process and its children), peak memory and the
    work folder's size to work/compute_log.jsonl."""
    me = resource.getrusage(resource.RUSAGE_SELF)
    kids = resource.getrusage(resource.RUSAGE_CHILDREN)
    cpu = me.ru_utime + me.ru_stime + kids.ru_utime + kids.ru_stime - cpu0
    # ru_maxrss is bytes on macOS and kilobytes on Linux.
    scale = 1 if sys.platform == "darwin" else 1024
    peak = max(me.ru_maxrss, kids.ru_maxrss) * scale
    disk = sum(p.stat().st_size for p in work.rglob("*") if p.is_file())
    row = {
        "step": step,
        "elapsed_s": round(time.time() - t0, 3),
        "cpu_s": round(cpu, 3),
        "peak_rss_bytes": peak,
        "work_dir_bytes": disk,
        "host": os.uname().nodename,
        "when": time.strftime("%Y-%m-%dT%H:%M:%S"),
        **(extra or {}),
    }
    with open(work / "compute_log.jsonl", "a") as f:
        f.write(json.dumps(row) + "\n")
    print(
        f"[{step}] {row['elapsed_s']} s elapsed, {row['cpu_s']} s CPU, "
        f"peak {peak / 1e6:.0f} MB, work folder {disk / 1e6:.1f} MB"
    )


def cpu_now() -> float:
    me = resource.getrusage(resource.RUSAGE_SELF)
    kids = resource.getrusage(resource.RUSAGE_CHILDREN)
    return me.ru_utime + me.ru_stime + kids.ru_utime + kids.ru_stime


# ---- import-evidence ----------------------------------------------------------------
def download(url: str, out: Path, refresh: bool = False) -> Path:
    if out.exists() and not refresh:
        return out
    out.parent.mkdir(parents=True, exist_ok=True)
    with urllib.request.urlopen(url, timeout=300) as r:
        data = r.read()
    out.write_bytes(data)
    return out


def elm_header(text: str) -> dict:
    meta = {}
    for line in text.splitlines():
        if not line.startswith("#"):
            break
        k, _, v = line[1:].partition(":")
        meta[k.strip()] = v.strip()
    return meta


def read_elm_tsv(path: Path) -> pl.DataFrame:
    body = "\n".join(l for l in path.read_text().splitlines() if not l.startswith("#"))
    return pl.read_csv(
        io.StringIO(body), separator="\t", infer_schema_length=0, quote_char='"'
    )


def uniprot_versions(accessions: list[str], cache: Path) -> dict[str, tuple[str, str]]:
    """Current UniProt sequence version and sequence for each canonical accession.
    Isoform accessions (P12345-2) are not entries of their own and are not looked up."""
    have = {}
    if cache.exists():
        for row in csv.DictReader(open(cache), delimiter="\t"):
            have[row["Entry"]] = (row["Sequence version"], row["Sequence"])
    todo = sorted({a for a in accessions if "-" not in a} - set(have))
    for i in range(0, len(todo), 100):
        batch = todo[i : i + 100]
        q = urllib.parse.quote(" OR ".join(f"accession:{a}" for a in batch))
        with urllib.request.urlopen(UNIPROT_TSV + q, timeout=300) as r:
            text = r.read().decode()
        for row in csv.DictReader(io.StringIO(text), delimiter="\t"):
            have[row["Entry"]] = (row["Sequence version"], row["Sequence"])
    with open(cache, "w") as f:
        f.write("Entry\tSequence version\tSequence\n")
        for a, (v, s) in sorted(have.items()):
            f.write(f"{a}\t{v}\t{s}\n")
    return have


def sanborn_label(z: float | None, plan: dict) -> str:
    """Class of an activation tile, from the paper's own Z-test."""
    if z is None:
        return "not measured"
    cut = plan["evidence"]["sanborn_labels"]
    if z >= cut["active_min_z"]:
        return "active"
    if z >= cut["weak_min_z"]:
        return "weak"
    return "inactive"


def write_evidence(
    rows: list[dict], seq_by_md5: dict[str, str], out: Path, raw: Path | None
) -> dict:
    """Deduplicate evidence rows and write evidence.parquet/.csv, the protein FASTA the
    searches run against, and a manifest with checksums and counts.

    Duplicate evidence: the same stretch of the same sequence with the same measurement
    and outcome, reached through two records. The first is kept; merged_ids lists all.
    """
    ev = pl.DataFrame(rows, infer_schema_length=None)
    key = ["source", "protein_md5", "start0", "end0", "measurement", "outcome"]
    ev = ev.with_columns(pl.col("evidence_id").alias("_id"))
    dup = ev.group_by(key, maintain_order=True).agg(pl.col("_id").alias("merged_ids"))
    ev = (
        ev.unique(subset=key, keep="first", maintain_order=True)
        .join(dup, on=key, how="left")
        .with_columns(pl.col("merged_ids").list.join(";"))
        .drop("_id")
    )
    out.mkdir(parents=True, exist_ok=True)
    ev.write_parquet(out / "evidence.parquet")
    ev.write_csv(out / "evidence.csv")
    prot = {}
    for m, acc in ev.select("protein_md5", "protein_accession").iter_rows():
        prot.setdefault(m, acc)
    write_fasta(
        [(f"{m}|{acc}", seq_by_md5[m]) for m, acc in sorted(prot.items())],
        out / "evidence_proteins.fasta",
    )
    return {
        "built": time.strftime("%Y-%m-%d"),
        "inputs": (
            {
                str(p.relative_to(raw)): sha256(p)
                for p in sorted(raw.rglob("*"))
                if p.is_file()
            }
            if raw
            else {}
        ),
        "outputs": {
            p.name: sha256(p)
            for p in sorted(out.iterdir())
            if p.is_file() and p.name != "manifest.json"
        },
        "duplicates_merged": int(ev["merged_ids"].str.count_matches(";").sum()),
        "records": ev.height,
        "proteins": len(prot),
        "by_source_and_outcome": ev.group_by("source", "outcome")
        .len()
        .sort("source", "outcome")
        .to_dicts(),
        "sequence_matches_current_uniprot": ev.group_by(
            "source", "sequence_matches_uniprot"
        )
        .len()
        .sort("source")
        .to_dicts(),
    }


def cmd_import_evidence(a):
    t0, c0 = time.time(), cpu_now()
    plan = load_plan()
    raw = Path(a.raw)
    work = Path(a.work)
    work.mkdir(parents=True, exist_ok=True)
    rows = []

    # ELM: one row per instance, every InstanceLogic kept (true negatives are measured
    # negatives). Coordinates: ELM Start/End are 1-based inclusive on ELM's copy of the
    # sequence (instances.fasta), so start0 = Start - 1, end0 = End.
    elm_tsv = download(ELM_INSTANCES_URL, raw / "elm" / "instances.tsv")
    download(ELM_CLASSES_URL, raw / "elm" / "elms_index.tsv")
    elm_fa = download(ELM_FASTA_URL, raw / "elm" / "instances.fasta")
    meta = elm_header(elm_tsv.read_text())
    elm_seq = {h.split("|")[1]: s for h, s in read_fasta(elm_fa)}
    inst = read_elm_tsv(elm_tsv)
    n_inst = inst.height
    assert n_inst == int(meta["NumInstances"]), (n_inst, meta)
    classes = read_elm_tsv(raw / "elm" / "elms_index.tsv")
    cls_desc = dict(zip(classes["ELMIdentifier"], classes["Description"]))

    # Sanborn et al. 2021, Figure 1 source data 1, sheets "TF sequences" and "TF tiles".
    import openpyxl

    xlsx = download(
        SANBORN_TILES_URL, raw / "sanborn2021" / "elife-68068-fig1-data1-v3.xlsx"
    )
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    tf_seq = {
        r[1]: r[2]
        for r in list(wb["TF sequences, see Fig 1"].iter_rows(values_only=True))[1:]
    }
    tiles_all = list(wb["TF tiles, see Fig 1"].iter_rows(values_only=True))
    head = list(tiles_all[0])
    assert head[:8] == [
        None,
        "protein sequence",
        "protein",
        "protein ID",
        "position",
        "activation",
        "Z score",
        "num cells",
    ], head
    tiles = [r for r in tiles_all[1:] if r[0] and r[0].startswith("A_tiles_")]

    accs = sorted(set(elm_seq) | set(tf_seq))
    uni = uniprot_versions(accs, raw / "uniprot_sequence_versions.tsv")

    skipped = Counter()
    for r in inst.iter_rows(named=True):
        acc = r["Primary_Acc"]
        seq = elm_seq.get(acc)
        if seq is None:
            skipped["ELM: no sequence in instances.fasta"] += 1
            continue
        s0, e0 = int(r["Start"]) - 1, int(r["End"])
        if not (0 <= s0 < e0 <= len(seq)):
            skipped["ELM: coordinates outside ELM's sequence"] += 1
            continue
        u = uni.get(acc)
        rows.append(
            dict(
                evidence_id=f"ELM:{r['Accession']}",
                source="ELM",
                source_version=f"ELM instance download {meta['ELM_Instance_Download_Version']}, "
                f"{meta['ELM_Instance_Download_Date'][:10]}",
                source_record_id=r["Accession"],
                source_url=f"http://elm.eu.org/instances/{r['Accession']}/",
                protein_accession=acc,
                protein_name=r["ProteinName"],
                organism=r["Organism"],
                sequence_origin="ELM instances.fasta",
                protein_md5=md5(seq),
                uniprot_sequence_version=u[0] if u else None,
                sequence_matches_uniprot=(u[1] == seq) if u else None,
                orig_start=int(r["Start"]),
                orig_end=int(r["End"]),
                orig_coordinates="1-based, inclusive (ELM)",
                start0=s0,
                end0=e0,
                stretch=seq[s0:e0],
                measurement=f"{r['ELMIdentifier']}: {cls_desc.get(r['ELMIdentifier'], '')}",
                outcome=r["InstanceLogic"],
                value=None,
                value_units=None,
                z_score=None,
                n_cells=None,
                assay_context=f"ELM-curated literature evidence; methods: {r['Methods']}",
                references="; ".join(f"PMID:{p}" for p in r["References"].split()),
                endpoint_note="ELM instance span is the curated motif, as the curators drew it.",
            )
        )
    for r in tiles:
        name, tseq, prot, acc, pos, act, z, ncells = r[:8]
        seq = tf_seq.get(acc)
        pos = int(pos)
        if seq is None or seq[pos : pos + len(tseq)] != tseq:
            skipped[
                "Sanborn: tile not found at its position in the paper's TF sequence"
            ] += 1
            continue
        u = uni.get(acc)
        z = None if z is None else float(z)
        rows.append(
            dict(
                evidence_id=f"SANBORN2021:{name}",
                source="Sanborn2021",
                source_version="eLife 10:e68068 (2021), Figure 1 source data 1, v3",
                source_record_id=name,
                source_url=f"https://doi.org/{SANBORN_DOI}",
                protein_accession=acc,
                protein_name=prot,
                organism="Saccharomyces cerevisiae",
                sequence_origin="Sanborn 2021 sheet 'TF sequences, see Fig 1'",
                protein_md5=md5(seq),
                uniprot_sequence_version=u[0] if u else None,
                sequence_matches_uniprot=(u[1] == seq) if u else None,
                orig_start=pos,
                orig_end=pos + len(tseq),
                orig_coordinates="0-based start, end exclusive (paper's 'position' + 53)",
                start0=pos,
                end0=pos + len(tseq),
                stretch=tseq,
                measurement="transcriptional activation of a 53-aa tile",
                outcome=sanborn_label(z, plan),
                value=None if act is None else float(act),
                value_units="fold increase in GFP reporter over background",
                z_score=z,
                n_cells=None if ncells is None else float(ncells),
                assay_context=(
                    "in vivo, S. cerevisiae; tile fused to an estrogen-inducible "
                    "artificial TF (mouse DNA-binding domain, mCherry tag) driving a "
                    "genome-integrated GFP reporter; FACS into 8 GFP bins + sequencing"
                ),
                references=f"PMID:{SANBORN_PMID}; doi:{SANBORN_DOI}",
                endpoint_note=(
                    "Tile ends were set by the library design (53 aa, step 12-13 aa); "
                    "they are not the ends of the residues the activation needs."
                ),
            )
        )

    seq_by_md5 = {md5(x): x for x in list(elm_seq.values()) + list(tf_seq.values())}
    manifest = write_evidence(rows, seq_by_md5, work / "evidence", raw)
    manifest.update(
        {
            "elm_instances_in_file": n_inst,
            "sanborn_tf_tiles_in_file": len(tiles),
            "sanborn_paper_says": 7460,
            "skipped": dict(skipped),
        }
    )
    (work / "evidence" / "manifest.json").write_text(
        json.dumps(manifest, indent=1, default=str)
    )
    print(
        json.dumps(
            {k: manifest[k] for k in manifest if k not in ("inputs", "outputs")},
            indent=1,
            default=str,
        )
    )
    log_compute(work, "import-evidence", t0, c0)


# ---- import-queries -----------------------------------------------------------------
QUEUE_COLUMNS = [
    "query_id",
    "received_at",
    "sequence",
    "source_accession",
    "source_start",
    "analyst_id",
    "investigation_note",
    "synthetic",
]


def cmd_import_queries(a):
    """Read the analyst's queue, drop what cannot be a query, and split by arrival order:
    the first n_dev usable queries are development, the next 100 are evaluation."""
    t0, c0 = time.time(), cpu_now()
    plan = load_plan()
    work = Path(a.work)
    q = pl.read_csv(a.queue, infer_schema_length=0)
    missing = [c for c in QUEUE_COLUMNS if c not in q.columns]
    if missing:
        raise SystemExit(
            f"queue file lacks columns {missing}; see templates/analyst_queue.csv"
        )
    q = q.with_columns(pl.col("sequence").str.to_uppercase().str.replace_all(r"\s", ""))
    q = q.sort("received_at", "query_id")
    ev = pl.read_parquet(work / "evidence" / "evidence.parquet")
    prots = dict(
        (n.split("|")[0], s)
        for n, s in read_fasta(work / "evidence" / "evidence_proteins.fasta")
    )
    spans = defaultdict(list)
    for m, s0, e0 in ev.select("protein_md5", "start0", "end0").iter_rows():
        spans[m].append((s0, e0))
    mo = plan["cards"]["min_overlap"]

    def copied(s: str) -> bool:
        """The query occurs letter for letter in an evidence protein, over a measured stretch."""
        for m, ps in prots.items():
            i = ps.find(s)
            while i >= 0:
                if any(
                    overlap(i, i + len(s), s0, e0) >= min(mo, e0 - s0)
                    for s0, e0 in spans[m]
                ):
                    return True
                i = ps.find(s, i + 1)
        return False

    lo, hi = plan["queries"]["min_length"], plan["queries"]["max_length"]
    reasons = []
    for r in q.iter_rows(named=True):
        s = r["sequence"] or ""
        why = ""
        if not lo <= len(s) <= hi:
            why = f"length {len(s)} outside {lo}-{hi}"
        elif set(s) - AMINO:
            why = f"non-standard letters {''.join(sorted(set(s) - AMINO))}"
        elif copied(s):
            why = "copied from the evidence collection (occurs exactly over a measured stretch)"
        reasons.append(why)
    q = q.with_columns(pl.Series("excluded_because", reasons))
    usable = q.filter(pl.col("excluded_because") == "")
    n_dev, n_eval = plan["queries"]["n_development"], plan["queries"]["n_evaluation"]
    split = (
        ["development"] * n_dev
        + ["evaluation"] * n_eval
        + ["unused"] * max(0, usable.height - n_dev - n_eval)
    )[: usable.height]
    usable = usable.with_columns(pl.Series("split", split))
    out = work / "queries"
    out.mkdir(parents=True, exist_ok=True)
    q.filter(pl.col("excluded_because") != "").write_csv(out / "excluded.csv")
    usable.write_csv(out / "queries.csv")
    for part in ("development", "evaluation"):
        sub = usable.filter(pl.col("split") == part)
        write_fasta(zip(sub["query_id"], sub["sequence"]), out / f"{part}.fasta")
    n = Counter(split)
    print(
        f"queue rows {q.height}, excluded {q.height - usable.height}, "
        f"development {n['development']}, evaluation {n['evaluation']}"
    )
    if n["evaluation"] < n_eval:
        print(f"only {n['evaluation']} evaluation queries; the plan needs {n_eval}")
    log_compute(work, "import-queries", t0, c0)


# ---- make-controls ------------------------------------------------------------------
def dipeptide_shuffle(seq: str, rng: random.Random) -> str:
    """A random sequence with the same first residue, last residue and count of every
    adjacent residue pair as `seq` (Altschul & Erickson 1985, Mol Biol Evol 2:526).

    Copied unchanged from nextflow-runs/shared/decoys.py on PR 138 (not merged into main
    when this was written); replace with an import once it is.
    """
    if len(seq) < 3:
        return seq
    out = defaultdict(list)
    for x, y in zip(seq, seq[1:]):
        out[x].append(y)
    last = seq[-1]
    while True:
        final = {v: rng.choice(ends) for v, ends in out.items() if v != last}
        ok = True
        for v in final:
            seen, u = set(), v
            while u != last:
                if u in seen:
                    ok = False
                    break
                seen.add(u)
                u = final[u]
            if not ok:
                break
        if ok:
            break
    order = {}
    for v, ends in out.items():
        rest = list(ends)
        if v in final:
            rest.remove(final[v])
        rng.shuffle(rest)
        order[v] = rest + ([final[v]] if v in final else [])
    walk, u = [seq[0]], seq[0]
    for _ in range(len(seq) - 1):
        u = order[u].pop(0)
        walk.append(u)
    return "".join(walk)


def composition_shuffle(seq: str, rng: random.Random) -> str:
    s = list(seq)
    rng.shuffle(s)
    return "".join(s)


def distinct_shuffles(seq, fn, n, rng, max_tries):
    """Up to n shuffles that differ from seq and from each other. Fewer than n means the
    sequence has too little variety for n distinct shuffles of this kind."""
    got, tries = [], 0
    while len(got) < n and tries < max_tries:
        tries += 1
        s = fn(seq, rng)
        if s != seq and s not in got:
            got.append(s)
    return got


def cmd_make_controls(a):
    t0, c0 = time.time(), cpu_now()
    plan = load_plan()
    work = Path(a.work)
    c = plan["controls"]
    q = pl.read_csv(work / "queries" / "queries.csv", infer_schema_length=0)
    q = q.filter(pl.col("split").is_in(["development", "evaluation"]))
    rows, recs = [], []
    for r in q.iter_rows(named=True):
        rng = random.Random(f"{c['seed']}:{r['query_id']}")
        for kind, fn in (
            ("composition", composition_shuffle),
            ("dipeptide", dipeptide_shuffle),
        ):
            got = distinct_shuffles(
                r["sequence"], fn, c["n_per_kind"], rng, c["max_tries"]
            )
            rows.append(
                dict(
                    query_id=r["query_id"],
                    split=r["split"],
                    kind=kind,
                    n_distinct=len(got),
                    low_diversity=len(got) < c["n_per_kind"],
                )
            )
            for i, s in enumerate(got, 1):
                recs.append((f"{r['query_id']}__{kind}{i:02d}", s))
    out = work / "controls"
    out.mkdir(parents=True, exist_ok=True)
    write_fasta(recs, out / "controls.fasta")
    flags = pl.DataFrame(rows)
    flags.write_csv(out / "control_diversity.csv")
    print(
        f"{len(recs)} control sequences; low-diversity flags: "
        f"{flags.filter(pl.col('low_diversity')).height} of {flags.height} query x kind"
    )
    log_compute(work, "make-controls", t0, c0)


# ---- hits ---------------------------------------------------------------------------
HIT_COLUMNS = [
    "workflow",
    "setting",
    "method",
    "query_id",
    "target",
    "q_start",
    "q_end",
    "t_start",
    "t_end",
    "q_aln",
    "t_aln",
    "score",
    "score_name",
    "evalue",
    "evalue_name",
]


# Workflow B is one setting: an exact full-length match, else BLAST with the task picked by
# query length. An exact match outranks every BLAST alignment.
B_SETTING = "exact+blast"
EXACT_SCORE = 1e6


def hit_schema():
    return {
        c: (
            pl.Float64
            if c in ("score", "evalue")
            else pl.Int64 if c in ("q_start", "q_end", "t_start", "t_end") else pl.Utf8
        )
        for c in HIT_COLUMNS
    }


def parent_query(qid: str) -> str:
    return qid.split("__")[0]


def split_fasta_by_len(records, cutoff):
    short = [(n, s) for n, s in records if len(s) < cutoff]
    long_ = [(n, s) for n, s in records if len(s) >= cutoff]
    return short, long_


def cmd_search_b(a):
    """Workflow B: exact full-alphabet substring lookup, then local BLAST with the task
    picked by query length (blastp-short below 30 residues, as the BLAST manual advises).
    """
    t0, c0 = time.time(), cpu_now()
    plan = load_plan()
    b = plan["workflows"]["B"]
    work = Path(a.work)
    bindir = Path(a.blast_bin)
    out = work / "hits"
    out.mkdir(parents=True, exist_ok=True)
    targets = read_fasta(work / "evidence" / "evidence_proteins.fasta")
    queries = []
    for part in ("development", "evaluation"):
        p = work / "queries" / f"{part}.fasta"
        if p.exists():
            queries += read_fasta(p)
    if (work / "controls" / "controls.fasta").exists():
        queries += read_fasta(work / "controls" / "controls.fasta")

    rows = []
    for qn, qs in queries:
        for tn, ts in targets:
            i = ts.find(qs)
            while i >= 0:
                rows.append(
                    dict(
                        workflow="B",
                        setting=B_SETTING,
                        method="exact",
                        query_id=qn,
                        target=tn,
                        q_start=0,
                        q_end=len(qs),
                        t_start=i,
                        t_end=i + len(qs),
                        q_aln=qs,
                        t_aln=ts[i : i + len(qs)],
                        score=EXACT_SCORE,
                        score_name="exact full-length match",
                        evalue=None,
                        evalue_name="none (exact match; no statistic)",
                    )
                )
                i = ts.find(qs, i + 1)

    db = work / "blastdb" / "evidence"
    db.parent.mkdir(exist_ok=True)
    tb = time.time()
    subprocess.run(
        [
            str(bindir / "makeblastdb"),
            "-in",
            str(work / "evidence" / "evidence_proteins.fasta"),
            "-dbtype",
            "prot",
            "-out",
            str(db),
            "-parse_seqids",
        ],
        check=True,
        capture_output=True,
    )
    index_s = time.time() - tb
    cols = [
        "qseqid",
        "sseqid",
        "qstart",
        "qend",
        "sstart",
        "send",
        "qseq",
        "sseq",
        "bitscore",
        "evalue",
    ]
    fmt = "6 " + " ".join(cols)
    short, long_ = split_fasta_by_len(queries, b["short_task_below"])
    parts = [pl.DataFrame(rows, schema=hit_schema())]
    for task, recs in (("blastp-short", short), ("blastp", long_)):
        if not recs:
            continue
        qf = out / f"b_{task}.query.fasta"
        tsv = out / f"b_{task}.tsv"
        write_fasta(recs, qf)
        cmd = [
            str(bindir / "blastp"),
            "-task",
            task,
            "-query",
            str(qf),
            "-db",
            str(db),
            "-outfmt",
            fmt,
            "-num_threads",
            str(a.threads),
            "-out",
            str(tsv),
        ]
        for k, v in b["tasks"][task].items():
            cmd += [f"-{k}", str(v)]
        (out / f"b_{task}.command.txt").write_text(" ".join(cmd) + "\n")
        subprocess.run(cmd, check=True)
        if tsv.stat().st_size == 0:
            continue
        # BLAST pads small numbers with spaces (" 9.1"), so read text and strip.
        t = pl.read_csv(
            tsv, separator="\t", has_header=False, new_columns=cols, infer_schema=False
        )
        t = t.with_columns(
            pl.col(c).str.strip_chars().cast(pl.Int64)
            for c in ("qstart", "qend", "sstart", "send")
        )
        t = t.with_columns(
            pl.col(c).str.strip_chars().cast(pl.Float64) for c in ("bitscore", "evalue")
        )
        parts.append(
            t.select(
                pl.lit("B").alias("workflow"),
                pl.lit(B_SETTING).alias("setting"),
                pl.lit(task).alias("method"),
                pl.col("qseqid").alias("query_id"),
                pl.col("sseqid").alias("target"),
                (pl.col("qstart") - 1).alias("q_start"),
                pl.col("qend").alias("q_end"),
                (pl.col("sstart") - 1).alias("t_start"),
                pl.col("send").alias("t_end"),
                pl.col("qseq").alias("q_aln"),
                pl.col("sseq").alias("t_aln"),
                pl.col("bitscore").cast(pl.Float64).alias("score"),
                pl.lit("BLAST bit score").alias("score_name"),
                pl.col("evalue").cast(pl.Float64),
                pl.lit(
                    f"BLAST {task} E-value (database size and alignment choice included; "
                    "not the choice of workflow)"
                ).alias("evalue_name"),
            ).cast(hit_schema())
        )
        tsv.unlink()
    hits = pl.concat(parts, how="vertical")
    hits.write_parquet(out / "hits_B.parquet")
    print(
        f"workflow B: {hits.height} hits for {len(queries)} sequences "
        f"({len(short)} with blastp-short)"
    )
    log_compute(
        work,
        "search-b",
        t0,
        c0,
        {"blast_index_build_s": round(index_s, 3), "threads": a.threads},
    )


def cmd_kmerseek_jobs(a):
    """Write work/sherlock_in/: the settings for the Sherlock array job (one line per
    alphabet and k), every query and shuffle in one FASTA, and the evidence proteins.
    kmerseek is not run here: on this laptop it is refused by a hook; it runs on Sherlock.
    """
    plan = load_plan()
    work = Path(a.work)
    grid = (
        plan["workflows"]["A"]["grid"]
        if not a.frozen
        else [load_lock(work)["choice"]["A"]]
    )
    lines = ["setting\talphabet\tksize"]
    for g in grid:
        for k in g.get("ksizes", [g.get("ksize")]):
            lines.append(f"{g['alphabet']}_k{k}\t{g['alphabet']}\t{k}")
    out = work / "sherlock_in"
    out.mkdir(exist_ok=True)
    (out / "kmerseek_settings.tsv").write_text("\n".join(lines) + "\n")
    recs = []
    for part in ("development", "evaluation"):
        if (work / "queries" / f"{part}.fasta").exists():
            recs += read_fasta(work / "queries" / f"{part}.fasta")
    if (work / "controls" / "controls.fasta").exists():
        recs += read_fasta(work / "controls" / "controls.fasta")
    write_fasta(recs, out / "queries_and_controls.fasta")
    shutil.copy(
        work / "evidence" / "evidence_proteins.fasta", out / "evidence_proteins.fasta"
    )
    print(
        f"{len(lines) - 1} kmerseek settings, {len(recs)} query and shuffle sequences -> {out}"
    )


def cmd_read_kmerseek(a):
    """Read kmerseek search CSVs (one per setting, named <setting>.csv) into the hit table.
    Score: region_ka_bits (extension on, as the plan requires); E-value: region_evalue,
    labelled as kmerseek's own and not calibrated for this use (notebook 344)."""
    work = Path(a.work)
    rows = []
    for p in sorted(Path(a.dir).glob("*.csv")):
        setting = p.stem
        df = pl.read_csv(p, infer_schema_length=10000)
        for r in df.iter_rows(named=True):
            rows.append(
                dict(
                    workflow="A",
                    setting=setting,
                    method="kmerseek",
                    query_id=r["query_name"].split()[0],
                    target=r["target_name"].split()[0],
                    q_start=int(r["region_start"]),
                    q_end=int(r["region_end"]),
                    t_start=int(r["target_start"]),
                    t_end=int(r["target_end"]),
                    q_aln=r["region_subseq"],
                    t_aln=r["target_subseq"],
                    score=float(r["region_ka_bits"]),
                    score_name="kmerseek region_ka_bits",
                    evalue=float(r["region_evalue"]),
                    evalue_name="kmerseek region_evalue (not calibrated for this use; "
                    "see notebook 344)",
                )
            )
    hits = pl.DataFrame(rows, schema=hit_schema())
    (work / "hits").mkdir(exist_ok=True)
    hits.write_parquet(work / "hits" / "hits_A.parquet")
    print(
        f"workflow A: {hits.height} region rows from {len(list(Path(a.dir).glob('*.csv')))} settings"
    )


# ---- cards --------------------------------------------------------------------------
def overlap(a0, a1, b0, b1):
    return max(0, min(a1, b1) - max(a0, b0))


def hits_to_cards(
    hits: pl.DataFrame, ev: pl.DataFrame, queries: pl.DataFrame, plan: dict
) -> pl.DataFrame:
    """One card per (query, workflow, setting, evidence record): the best hit on that record.

    A hit supports a record when the hit's target span overlaps the record's stretch by at
    least min(min_overlap, record length) residues. Self-records are removed: a record on
    the protein the query was taken from (same accession) whose stretch overlaps the
    query's own position there."""
    mo = plan["cards"]["min_overlap"]
    by_prot = defaultdict(list)
    for r in ev.select(
        "evidence_id", "protein_md5", "protein_accession", "start0", "end0"
    ).iter_rows():
        by_prot[r[1]].append(r)
    qsrc = {r["query_id"]: r for r in queries.iter_rows(named=True)}
    best = {}
    for h in hits.iter_rows(named=True):
        pm = h["target"].split("|")[0]
        src = qsrc.get(parent_query(h["query_id"]), {})
        for eid, _, acc, s0, e0 in by_prot.get(pm, []):
            ov = overlap(h["t_start"], h["t_end"], s0, e0)
            if ov < min(mo, e0 - s0):
                continue
            if (
                src.get("source_accession")
                and acc.split("-")[0] == src["source_accession"].split("-")[0]
            ):
                qs = int(src["source_start"] or -(10**9))
                if overlap(qs, qs + len(src["sequence"]), s0, e0) > 0:
                    continue
            k = (h["query_id"], h["workflow"], h["setting"], eid)
            if k not in best or h["score"] > best[k]["score"]:
                best[k] = {**h, "evidence_id": eid, "overlap": ov}
    if not best:
        return pl.DataFrame(
            schema={**{c: pl.Utf8 for c in HIT_COLUMNS}, "evidence_id": pl.Utf8}
        )
    return pl.DataFrame(list(best.values()), infer_schema_length=None)


def shuffle_expectation(
    cards: pl.DataFrame, plan: dict
) -> tuple[pl.DataFrame, pl.DataFrame]:
    """For each card: how many cards at least this good one shuffled copy of the same query
    produces, on average over that query's shuffles of the plan's kind. Shuffles go through
    the same search, region choice and evidence mapping as the real query, so the number
    counts every choice made after the query is given, against this database.

    A real card is compared with all n shuffles. A shuffle's own card is compared with the
    other n - 1 shuffles of its query, so the same cut-off applied to shuffles estimates how
    many cards a query with no real relative would keep."""
    kind = plan["controls"]["kind_for_expectation"]
    n_shuf = plan["controls"]["n_per_kind"]
    cards = cards.with_columns(
        pl.col("query_id")
        .map_elements(parent_query, return_dtype=pl.Utf8)
        .alias("parent"),
        pl.col("query_id").str.contains("__").alias("is_control"),
    )
    ctrl = cards.filter(pl.col("is_control"))
    real = cards.filter(~pl.col("is_control"))
    cs = defaultdict(list)
    for r in (
        ctrl.filter(pl.col("query_id").str.contains(f"__{kind}"))
        .select("parent", "workflow", "setting", "query_id", "score")
        .iter_rows()
    ):
        cs[r[:3]].append((r[3], r[4]))
    exp = [
        sum(1 for _, s in cs.get((p, w, st), []) if s >= sc) / n_shuf
        for p, w, st, sc in real.select(
            "parent", "workflow", "setting", "score"
        ).iter_rows()
    ]
    loo = [
        sum(1 for q2, s in cs.get((p, w, st), []) if s >= sc and q2 != qid)
        / (n_shuf - 1)
        for p, w, st, qid, sc in ctrl.select(
            "parent", "workflow", "setting", "query_id", "score"
        ).iter_rows()
    ]
    return (
        real.with_columns(pl.Series("shuffle_expected_cards", exp, dtype=pl.Float64)),
        ctrl.with_columns(
            pl.Series("shuffle_expected_cards", loo, dtype=pl.Float64),
            pl.col("query_id").str.extract(r"__([a-z]+)\d+$").alias("control_kind"),
        ),
    )


def cmd_cards(a):
    t0, c0 = time.time(), cpu_now()
    plan = load_plan()
    work = Path(a.work)
    ev = pl.read_parquet(work / "evidence" / "evidence.parquet")
    q = pl.read_csv(work / "queries" / "queries.csv", infer_schema_length=0)
    parts = [pl.read_parquet(p) for p in sorted((work / "hits").glob("hits_*.parquet"))]
    hits = pl.concat([p for p in parts if p.height], how="vertical_relaxed")
    cards = hits_to_cards(hits, ev, q, plan)
    real, ctrl = shuffle_expectation(cards, plan)
    real = real.join(q.select("query_id", "split"), on="query_id", how="left")
    real.write_parquet(work / "hits" / "cards.parquet")
    ctrl.write_parquet(work / "hits" / "control_cards.parquet")
    print(f"{real.height} cards on real queries, {ctrl.height} on shuffles")
    log_compute(work, "cards", t0, c0)


# ---- tune ---------------------------------------------------------------------------
def kept(cards: pl.DataFrame, setting: str, max_shuffle_cards: float) -> pl.DataFrame:
    return cards.filter(
        (pl.col("setting") == setting)
        & (pl.col("shuffle_expected_cards") <= max_shuffle_cards)
    )


def cmd_tune(a):
    """Choose, on development queries only, workflow A's alphabet and k and each
    workflow's card cut-off. A card is kept when its shuffle-expected count is at most the
    cut-off, so the cut-off is itself the number of chance cards a query with no relative
    in the database is expected to keep. The rule (in the plan): among cut-offs up to
    `max_cutoff`, pick the setting and cut-off with the most development queries holding a
    card judged useful; ties go to the smaller k, then the alphabet listed first in the
    grid, then the stricter cut-off.

    control_cards_per_shuffle is reported as a check, not used to choose: the same rule
    applied to each shuffle against the other 19 should keep about the cut-off in cards
    per shuffle (at most one card per 20 shuffles more). Much more than that means
    the shuffles' scores are not exchangeable and the cut-off does not mean what it says.
    """
    plan = load_plan()
    t = plan["tuning"]
    kind = plan["controls"]["kind_for_expectation"]
    n_shuf = plan["controls"]["n_per_kind"]
    work = Path(a.work)
    cards = pl.read_parquet(work / "hits" / "cards.parquet").filter(
        pl.col("split") == "development"
    )
    q = pl.read_csv(work / "queries" / "queries.csv", infer_schema_length=0)
    dev_q = q.filter(pl.col("split") == "development")["query_id"].to_list()
    ctrl = pl.read_parquet(work / "hits" / "control_cards.parquet").filter(
        pl.col("parent").is_in(dev_q) & (pl.col("control_kind") == kind)
    )
    labels = pl.read_csv(a.dev_labels, infer_schema_length=0)
    useful = {
        (r["query_id"], r["evidence_id"])
        for r in labels.iter_rows(named=True)
        if (r["useful"] or "").strip().lower() == "yes" and (r["reason"] or "").strip()
    }
    order = {g["alphabet"]: i for i, g in enumerate(plan["workflows"]["A"]["grid"])}
    settings = sorted(
        set(cards.select("workflow", "setting").iter_rows())
        | set(ctrl.select("workflow", "setting").iter_rows())
    )
    table = []
    for wf, st in settings:
        for cut in t["shuffle_card_cutoffs"]:
            k = kept(cards, st, cut)
            n_useful = len(
                {
                    qid
                    for qid, eid in k.select("query_id", "evidence_id").iter_rows()
                    if (qid, eid) in useful
                }
            )
            ck = kept(ctrl, st, cut).height
            m = re.match(r"(.+)_k(\d+)$", st)
            table.append(
                dict(
                    workflow=wf,
                    setting=st,
                    cutoff=cut,
                    dev_useful_queries=n_useful,
                    control_cards_per_shuffle=round(
                        ck / max(1, len(dev_q) * n_shuf), 4
                    ),
                    alphabet_rank=order.get(m.group(1), 99) if m else 0,
                    ksize=int(m.group(2)) if m else 0,
                )
            )
    tab = pl.DataFrame(table)
    tab.write_csv(work / "tuning_table.csv")
    ok = tab.filter(pl.col("cutoff") <= t["max_cutoff"])
    choice = {}
    for wf in ("A", "B"):
        sub = ok.filter(pl.col("workflow") == wf).sort(
            ["dev_useful_queries", "ksize", "alphabet_rank", "cutoff"],
            descending=[True, False, False, False],
        )
        if sub.height == 0:
            choice[wf] = None
            continue
        r = sub.row(0, named=True)
        m = re.match(r"(.+)_k(\d+)$", r["setting"])
        choice[wf] = dict(
            setting=r["setting"],
            max_shuffle_cards=r["cutoff"],
            dev_useful_queries=r["dev_useful_queries"],
            control_cards_per_shuffle=r["control_cards_per_shuffle"],
            **({"alphabet": m.group(1), "ksize": int(m.group(2))} if m else {}),
        )
    (work / "tuning_choice.json").write_text(json.dumps(choice, indent=1))
    print(json.dumps(choice, indent=1))


# ---- freeze -------------------------------------------------------------------------
def load_lock(work: Path) -> dict:
    return json.loads((work / "evaluation_plan.lock.json").read_text())


LOCKED = [
    "code:evaluation_plan.yaml",
    "code:mr353.py",
    "work:evidence/manifest.json",
    "work:evidence/evidence.parquet",
    "work:evidence/evidence_proteins.fasta",
    "work:queries/queries.csv",
    "work:tuning_choice.json",
]


def locked_path(key: str, work: Path) -> Path:
    where, rel = key.split(":", 1)
    return (HERE if where == "code" else work) / rel


def cmd_freeze(a):
    """Record checksums of the plan, the code, the evidence collection, the queue and the
    tuning choice. `score` refuses to give PASS or FAIL if any changed afterwards."""
    work = Path(a.work)
    paths = {k: locked_path(k, work) for k in LOCKED}
    lock = {
        "frozen_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "checksums": {k: sha256(p) for k, p in paths.items() if p.exists()},
        "missing": [k for k, p in paths.items() if not p.exists()],
        "choice": (
            json.loads(paths["work:tuning_choice.json"].read_text())
            if paths["work:tuning_choice.json"].exists()
            else None
        ),
        "queue_file": a.queue,
        "queue_sha256": (
            sha256(Path(a.queue)) if a.queue and Path(a.queue).exists() else None
        ),
    }
    (work / "evaluation_plan.lock.json").write_text(json.dumps(lock, indent=1))
    print(json.dumps(lock, indent=1))


def check_lock(work: Path) -> list[str]:
    p = work / "evaluation_plan.lock.json"
    if not p.exists():
        return ["no evaluation_plan.lock.json: the plan was never frozen"]
    lock = load_lock(work)
    bad = [
        f"{k} changed after freezing"
        for k, h in lock["checksums"].items()
        if not locked_path(k, work).exists() or sha256(locked_path(k, work)) != h
    ]
    return bad + [f"{m} was missing when frozen" for m in lock["missing"]]


# ---- review-pack --------------------------------------------------------------------
def frozen_cards(work: Path, choice: dict) -> pl.DataFrame:
    cards = pl.read_parquet(work / "hits" / "cards.parquet").filter(
        pl.col("split") == "evaluation"
    )
    keep = []
    for wf in ("A", "B"):
        c = choice.get(wf)
        if c:
            keep.append(
                kept(
                    cards.filter(pl.col("workflow") == wf),
                    c["setting"],
                    c["max_shuffle_cards"],
                )
            )
    return pl.concat(keep, how="vertical_relaxed") if keep else cards.head(0)


def cmd_session_pack(a):
    """session_order.csv, for the analyst: which workflows to run on each evaluation query
    and in which order. Orders cycle through every permutation and are dealt to queries at
    random, so each order is used equally often. Each row after the first for a query is a
    second look at that query (prior_exposure = yes)."""
    import itertools

    plan = load_plan()
    work = Path(a.work)
    rng = random.Random(plan["review"]["seed"])
    q = pl.read_csv(work / "queries" / "queries.csv", infer_schema_length=0).filter(
        pl.col("split") == "evaluation"
    )
    wfs = ["A", "B"] + (["C"] if a.with_c else [])
    # Cyclic rotations first, then rotations of the reversed order: when the queries do not
    # divide evenly, the orders used once more are rotations of each other, so every
    # workflow still comes first within one query of the others.
    rot = [wfs[i:] + wfs[:i] for i in range(len(wfs))]
    rev = [w[::-1] for w in rot]
    perms = [tuple(p) for p in rot] + [
        tuple(p) for p in rev if tuple(p) not in map(tuple, rot)
    ]
    assert sorted(perms) == sorted(itertools.permutations(wfs)), perms
    deal = [perms[i % len(perms)] for i in range(q.height)]
    rng.shuffle(deal)
    out = work / "review"
    out.mkdir(exist_ok=True)
    with open(out / "session_order.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(
            [
                "query_id",
                "order_position",
                "workflow",
                "prior_exposure",
                "analyst_id",
                "search_seconds",
                "reading_seconds",
                "verification_seconds",
                "retained_evidence_ids",
                "notes",
            ]
        )
        for qid, perm in zip(q["query_id"], deal):
            for i, wf in enumerate(perm, 1):
                w.writerow(
                    [qid, i, wf, "yes" if i > 1 else "no", "", "", "", "", "", ""]
                )
    print(f"{q.height} evaluation queries, orders {dict(Counter(deal))}")


def cmd_review_pack(a):
    """blind_review.csv, for reviewers: every card workflows A and B returned under the
    frozen settings, plus every measurement the analyst kept with workflow C (an evidence
    id, or a free citation), deduplicated per query, shuffled, with no workflow column.
    blind_key.csv says which workflow(s) returned each card; reviewers do not see it."""
    plan = load_plan()
    work = Path(a.work)
    lock = load_lock(work)
    rng = random.Random(plan["review"]["seed"] + 1)
    ev = pl.read_parquet(work / "evidence" / "evidence.parquet")
    evd = {r["evidence_id"]: r for r in ev.iter_rows(named=True)}
    cards = frozen_cards(work, lock["choice"])
    pool = defaultdict(set)
    for qid, eid, wf in cards.select("query_id", "evidence_id", "workflow").iter_rows():
        pool[(qid, eid)].add(wf)
    if a.session:
        for r in csv.DictReader(open(a.session)):
            if r["workflow"] == "C":
                for item in filter(
                    None,
                    (x.strip() for x in (r["retained_evidence_ids"] or "").split(";")),
                ):
                    pool[(r["query_id"], item)].add("C")
    items = sorted(pool.items())
    rng.shuffle(items)
    out = work / "review"
    out.mkdir(exist_ok=True)
    with open(out / "blind_review.csv", "w", newline="") as f, open(
        out / "blind_key.csv", "w", newline=""
    ) as g:
        w, k = csv.writer(f), csv.writer(g)
        w.writerow(
            [
                "card_id",
                "query_id",
                "evidence_id",
                "protein",
                "stretch",
                "outcome",
                "value",
                "value_units",
                "assay_context",
                "references",
                "source_url",
                "previously_seen_by_analyst",
                "useful",
                "reason",
                "reviewer_id",
            ]
        )
        k.writerow(["card_id", "workflows"])
        for i, ((qid, eid), wset) in enumerate(items, 1):
            e = evd.get(eid)
            cid = f"card{i:05d}"
            if (
                e is None
            ):  # a citation the analyst kept with workflow C, outside this collection
                w.writerow(
                    [cid, qid, "", "", "", "", "", "", "", eid, "", "", "", "", ""]
                )
            else:
                w.writerow(
                    [
                        cid,
                        qid,
                        eid,
                        f"{e['protein_accession']} {e['protein_name']}",
                        e["stretch"],
                        e["outcome"],
                        e["value"],
                        e["value_units"],
                        e["assay_context"],
                        e["references"],
                        e["source_url"],
                        "",
                        "",
                        "",
                        "",
                    ]
                )
            k.writerow([cid, ";".join(sorted(wset))])
    print(f"{len(items)} cards to review")


# ---- score --------------------------------------------------------------------------
def wilson(k, n, z=1.96):
    if n == 0:
        return (float("nan"), float("nan"))
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return (round(c - h, 4), round(c + h, 4))


def bootstrap_median(x, n, seed):
    rng = random.Random(seed)
    if not x:
        return (float("nan"), float("nan"))
    meds = sorted(statistics.median(rng.choices(x, k=len(x))) for _ in range(n))
    return (round(meds[int(0.025 * n)], 4), round(meds[int(0.975 * n) - 1], 4))


def cmd_score(a):
    """Apply the frozen pass rule. Every reason the result cannot be read is listed; any
    one of them makes the verdict INSUFFICIENT_EVIDENCE."""
    plan = load_plan()
    rule = plan["pass_rule"]
    work = Path(a.work)
    out = work / "results"
    out.mkdir(exist_ok=True)
    reasons = []
    reasons += check_lock(work)
    q = pl.read_csv(work / "queries" / "queries.csv", infer_schema_length=0)
    ev_q = q.filter(pl.col("split") == "evaluation")
    if ev_q.height < plan["queries"]["n_evaluation"]:
        reasons.append(
            f"{ev_q.height} evaluation queries; the plan needs {plan['queries']['n_evaluation']}"
        )
    if (ev_q["synthetic"].str.to_lowercase() == "yes").any():
        reasons.append(
            "evaluation queries are synthetic fixtures, not an analyst's queue"
        )
    review = Path(a.review) if a.review else None
    session = Path(a.session) if a.session else None
    key = work / "review" / "blind_key.csv"
    if not (review and review.exists()):
        reasons.append("no completed blind review sheet")
    if not (session and session.exists()):
        reasons.append("no completed timed-session sheet")

    per_query = []
    verdict_detail = {}
    if review and review.exists() and session and session.exists() and key.exists():
        rv = pl.read_csv(review, infer_schema_length=0)
        if (
            (rv["synthetic"].str.to_lowercase() == "yes").any()
            if "synthetic" in rv.columns
            else False
        ):
            reasons.append("review judgments are synthetic fixtures")
        kk = dict(csv.reader(open(key)))
        useful_by = defaultdict(set)
        bad_rows = 0
        for r in rv.iter_rows(named=True):
            u = (r["useful"] or "").strip().lower()
            if u not in ("yes", "no"):
                bad_rows += 1
                continue
            seen = (r["previously_seen_by_analyst"] or "").strip().lower()
            if u == "yes" and (
                not (r["reason"] or "").strip() or seen not in ("yes", "no")
            ):
                bad_rows += 1
                continue
            if u == "yes" and seen == "no":
                for wf in kk.get(r["card_id"], "").split(";"):
                    useful_by[wf].add(r["query_id"])
        if bad_rows:
            reasons.append(
                f"{bad_rows} review rows lack a yes/no judgment, the written reason, "
                "or whether the analyst had seen the measurement before"
            )
        ss = pl.read_csv(session, infer_schema_length=0)
        if (
            "synthetic" in ss.columns
            and (ss["synthetic"].str.to_lowercase() == "yes").any()
        ):
            reasons.append("timings are synthetic fixtures")
        tcol = ["search_seconds", "reading_seconds", "verification_seconds"]
        times, first = {}, {}
        missing_t = 0
        for r in ss.iter_rows(named=True):
            first[(r["query_id"], r["workflow"])] = r["order_position"] == "1"
            try:
                times[(r["query_id"], r["workflow"])] = sum(float(r[c]) for c in tcol)
            except (TypeError, ValueError):
                missing_t += 1
        if missing_t:
            reasons.append(
                f"{missing_t} timed-session rows lack one of the three times"
            )
        wfs = sorted({w for _, w in times})
        qids = ev_q["query_id"].to_list()
        yields = {w: sum(1 for x in qids if x in useful_by[w]) for w in wfs}
        comps = [w for w in wfs if w != "A"]
        strongest = (
            max(
                comps,
                key=lambda w: (
                    yields[w],
                    -statistics.median(
                        [times[(x, w)] for x in qids if (x, w) in times]
                        or [float("inf")]
                    ),
                ),
            )
            if comps
            else None
        )
        red = []
        for x in qids:
            ta, tc = times.get((x, "A")), times.get((x, strongest))
            r = None if ta is None or not tc else (tc - ta) / tc
            if r is not None:
                red.append(r)
            per_query.append(
                dict(
                    query_id=x,
                    A_order=1 if first.get((x, "A")) else 2,
                    **{f"useful_{w}": x in useful_by[w] for w in wfs},
                    **{f"seconds_{w}": times.get((x, w)) for w in wfs},
                    reduction_vs_strongest=r,
                )
            )
        med = statistics.median(red) if red else float("nan")
        n = len(qids)
        c1 = yields.get("A", 0) >= rule["min_useful_queries"]
        c2 = strongest is None or yields.get("A", 0) >= yields[strongest]
        c3 = red != [] and med >= rule["min_median_time_reduction"]

        # Carry-over: a second look at a query may be faster because the analyst has seen it.
        def median_or_none(xs):
            return statistics.median(xs) if xs else None

        carryover = {
            w: dict(
                median_seconds_first_look=median_or_none(
                    [t for (x, ww), t in times.items() if ww == w and first[(x, ww)]]
                ),
                median_seconds_later_look=median_or_none(
                    [
                        t
                        for (x, ww), t in times.items()
                        if ww == w and not first[(x, ww)]
                    ]
                ),
                useful_first_look=sum(
                    1 for x in qids if x in useful_by[w] and first.get((x, w))
                ),
                useful_later_look=sum(
                    1 for x in qids if x in useful_by[w] and first.get((x, w)) is False
                ),
            )
            for w in wfs
        }
        a_first = [
            r
            for x, r in zip(
                [
                    p["query_id"]
                    for p in per_query
                    if p["reduction_vs_strongest"] is not None
                ],
                red,
            )
            if first.get((x, "A"))
        ]
        a_later = [
            r
            for x, r in zip(
                [
                    p["query_id"]
                    for p in per_query
                    if p["reduction_vs_strongest"] is not None
                ],
                red,
            )
            if first.get((x, "A")) is False
        ]
        verdict_detail = dict(
            carryover=carryover,
            median_reduction_when_A_first=median_or_none(a_first),
            median_reduction_when_A_later=median_or_none(a_later),
            yields={w: dict(n=yields[w], of=n, ci95=wilson(yields[w], n)) for w in wfs},
            strongest_comparison=strongest,
            median_reduction=med,
            median_reduction_ci95=bootstrap_median(
                red, plan["scoring"]["bootstrap"], plan["scoring"]["seed"]
            ),
            criteria={
                "1_at_least_20_useful": c1,
                "2_yield_not_below_strongest": c2,
                "3_median_reduction_at_least_50pct": c3,
            },
        )
    pl.DataFrame(per_query, infer_schema_length=None).write_csv(
        out / "per_query_outcomes.csv"
    )
    if reasons:
        verdict = "INSUFFICIENT_EVIDENCE"
    elif all(verdict_detail["criteria"].values()):
        verdict = "PASS"
    else:
        verdict = "FAIL"
    result = dict(verdict=verdict, reasons=reasons, **verdict_detail)
    (out / "verdict.json").write_text(json.dumps(result, indent=1, default=str))
    print(json.dumps(result, indent=1, default=str))


# ---- viewer -------------------------------------------------------------------------
VIEWER = """<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Evidence cards</title><style>
:root{--bg:#fff;--fg:#1d1d1f;--mute:#5f6368;--line:#d6d6d6;--hit:#ffe08a;--card:#f7f7f8}
@media (prefers-color-scheme:dark){:root{--bg:#161618;--fg:#ececec;--mute:#a0a0a8;--line:#3a3a40;--hit:#7a5d00;--card:#212126}}
body{background:var(--bg);color:var(--fg);font:14px/1.45 -apple-system,Helvetica,Arial,sans-serif;margin:0 16px 40px;max-width:1100px}
h1{font-size:20px}.mute{color:var(--mute)}.q{border-top:2px solid var(--line);margin-top:22px;padding-top:8px}
.card{background:var(--card);border:1px solid var(--line);border-radius:6px;padding:8px 10px;margin:8px 0}
.seq{font-family:Menlo,"Courier New",monospace;font-size:13px;word-break:break-all}
mark{background:var(--hit);color:inherit}table{border-collapse:collapse}td{padding:1px 8px 1px 0;vertical-align:top}
label{margin-right:12px}
</style></head><body>
<h1>Evidence cards</h1>
<p class="mute">__BANNER__</p>
<p>How to read a card: the match line marks '|' where query and target have the same
amino acid, '.' where they differ and a blank at a gap. The highlighted letters are the measured stretch's residues that the
match covers. The match says the two stretches share sequence; it does not say they have the
same biological effect. Tile ends (Sanborn 2021) come from the library design, not from the
residues the activation needs.</p>
<p><label><input type="checkbox" id="blind"> hide which workflow found each card</label>
<label>query <select id="pick"></select></label></p>
<div id="out"></div>
<script>
const DATA=__DATA__;
const esc=s=>String(s??'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
// '|' same letter, '.' different letter, ' ' gap
function matchLine(q,t){let m='';for(let i=0;i<q.length;i++){m+=(q[i]==='-'||t[i]==='-')?' ':(q[i]===t[i]?'|':'.');}return m.replace(/ /g,'&nbsp;');}
function matchCount(q,t){let n=0;for(let i=0;i<q.length;i++){if(q[i]!=='-'&&q[i]===t[i])n++;}return `${n} of ${q.length} identical`;}
function mark(s,a,b){return esc(s.slice(0,a))+'<mark>'+esc(s.slice(a,b))+'</mark>'+esc(s.slice(b));}
const pick=document.getElementById('pick');
['(all)',...Object.keys(DATA.queries)].forEach(q=>{const o=document.createElement('option');o.textContent=q;pick.appendChild(o)});
function render(){
 const blind=document.getElementById('blind').checked, sel=pick.value; let h='';
 for(const [qid,q] of Object.entries(DATA.queries)){
  if(sel!=='(all)'&&sel!==qid) continue;
  h+=`<div class="q"><b>${esc(qid)}</b> <span class="mute">${esc(q.split)} · ${q.seq.length} aa</span><div class="seq">${esc(q.seq)}</div>`;
  if(!q.cards.length) h+='<p class="mute">no card passed the frozen cut-off</p>';
  for(const c of q.cards){
   const e=DATA.evidence[c.evidence_id];
   const a=Math.max(c.t_start,e.start0)-e.start0, b=Math.min(c.t_end,e.end0)-e.start0;
   h+=`<div class="card"><table>
   ${blind?'':`<tr><td class="mute">found by</td><td>${esc(c.workflows)}</td></tr>`}
   <tr><td class="mute">query ${c.q_start}-${c.q_end}</td><td class="seq">${esc(c.q_aln)}</td></tr>
   <tr><td class="mute">${matchCount(c.q_aln,c.t_aln)}</td><td class="seq">${matchLine(c.q_aln,c.t_aln)}</td></tr>
   <tr><td class="mute">target ${c.t_start}-${c.t_end}</td><td class="seq">${esc(c.t_aln)}</td></tr>
   <tr><td class="mute">measured stretch</td><td class="seq">${mark(e.stretch,a,b)} <span class="mute">(${e.start0}-${e.end0}, 0-based; source: ${esc(e.orig_coordinates)} ${e.orig_start}-${e.orig_end})</span></td></tr>
   <tr><td class="mute">protein</td><td>${esc(e.protein_accession)} ${esc(e.protein_name)} · ${esc(e.organism)} · ${e.uniprot_sequence_version!=null?`UniProt sequence version ${esc(e.uniprot_sequence_version)}`:'UniProt version not looked up'}${e.sequence_matches_uniprot===false?' (this sequence differs from current UniProt)':''}</td></tr>
   <tr><td class="mute">measurement</td><td>${esc(e.measurement)} · <b>${esc(e.outcome)}</b>${e.value!=null?` · ${Number(e.value).toFixed(2)} ${esc(e.value_units)}, Z ${Number(e.z_score).toFixed(2)}`:''}</td></tr>
   <tr><td class="mute">assay</td><td>${esc(e.assay_context)}</td></tr>
   <tr><td class="mute">source</td><td><a href="${esc(e.source_url)}">${esc(e.evidence_id)}</a> · ${esc(e.references)} · ${esc(e.source_version)}</td></tr>
   <tr><td class="mute">statistics</td><td>${esc(c.score_name)} ${Number(c.score).toFixed(1)}; expected cards this good from one shuffled query: ${Number(c.shuffle_expected_cards).toFixed(2)}${c.evalue!=null?`; ${esc(c.evalue_name)}: ${Number(c.evalue).toExponential(1)}`:''}</td></tr>
   <tr><td class="mute">note</td><td>${esc(e.endpoint_note)}</td></tr></table></div>`;
  }
  h+='</div>';
 }
 document.getElementById('out').innerHTML=h;
}
pick.onchange=render;document.getElementById('blind').onchange=render;render();
</script></body></html>"""


def cmd_viewer(a):
    work = Path(a.work)
    ev = pl.read_parquet(work / "evidence" / "evidence.parquet")
    q = pl.read_csv(work / "queries" / "queries.csv", infer_schema_length=0)
    q = q.filter(pl.col("split").is_in(["development", "evaluation"]))
    lockp = work / "evaluation_plan.lock.json"
    if lockp.exists() and load_lock(work)["choice"]:
        cards = frozen_cards(work, load_lock(work)["choice"])
        q = q.filter(pl.col("split") == "evaluation")
        banner = "Evaluation queries; cards kept by the frozen settings."
    else:
        cards = pl.read_parquet(work / "hits" / "cards.parquet")
        banner = "Not frozen: every card from every setting is shown."
    if a.banner:
        banner = a.banner + " " + banner
    merged = defaultdict(lambda: None)
    wf_by = defaultdict(set)
    for r in cards.sort("score", descending=True).iter_rows(named=True):
        k = (r["query_id"], r["evidence_id"])
        wf_by[k].add(f"{r['workflow']} ({r['setting']})")
        if merged[k] is None:
            merged[k] = r
    data = {"queries": {}, "evidence": {}}
    for r in q.iter_rows(named=True):
        data["queries"][r["query_id"]] = {
            "seq": r["sequence"],
            "split": r["split"],
            "cards": [],
        }
    need = set()
    for (qid, eid), r in merged.items():
        if qid in data["queries"]:
            data["queries"][qid]["cards"].append(
                {
                    **{
                        c: r[c]
                        for c in (
                            "evidence_id",
                            "q_start",
                            "q_end",
                            "t_start",
                            "t_end",
                            "q_aln",
                            "t_aln",
                            "score",
                            "score_name",
                            "evalue",
                            "evalue_name",
                            "shuffle_expected_cards",
                        )
                    },
                    "workflows": "; ".join(sorted(wf_by[(qid, eid)])),
                }
            )
            need.add(eid)
    for r in ev.filter(pl.col("evidence_id").is_in(list(need))).iter_rows(named=True):
        data["evidence"][r["evidence_id"]] = r
    page = VIEWER.replace(
        "__DATA__", json.dumps(data, default=str).replace("</", "<\\/")
    ).replace("__BANNER__", html.escape(banner))
    Path(a.out).write_text(page)
    print(
        f"viewer: {sum(len(v['cards']) for v in data['queries'].values())} cards -> {a.out}"
    )


# ---- main ---------------------------------------------------------------------------
def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    p.add_argument("--work", default="work", help="folder for every intermediate file")
    s = p.add_subparsers(dest="cmd", required=True)
    x = s.add_parser("import-evidence")
    x.add_argument("--raw", required=True)
    x.set_defaults(f=cmd_import_evidence)
    x = s.add_parser("import-queries")
    x.add_argument("--queue", required=True)
    x.set_defaults(f=cmd_import_queries)
    x = s.add_parser("make-controls")
    x.set_defaults(f=cmd_make_controls)
    x = s.add_parser("search-b")
    x.add_argument("--blast-bin", required=True)
    x.add_argument("--threads", type=int, default=4)
    x.set_defaults(f=cmd_search_b)
    x = s.add_parser("kmerseek-jobs")
    x.add_argument("--frozen", action="store_true")
    x.set_defaults(f=cmd_kmerseek_jobs)
    x = s.add_parser("read-kmerseek")
    x.add_argument("--dir", required=True)
    x.set_defaults(f=cmd_read_kmerseek)
    x = s.add_parser("cards")
    x.set_defaults(f=cmd_cards)
    x = s.add_parser("tune")
    x.add_argument("--dev-labels", required=True)
    x.set_defaults(f=cmd_tune)
    x = s.add_parser("freeze")
    x.add_argument("--queue")
    x.set_defaults(f=cmd_freeze)
    x = s.add_parser("session-pack")
    x.add_argument("--with-c", action="store_true")
    x.set_defaults(f=cmd_session_pack)
    x = s.add_parser("review-pack")
    x.add_argument("--session", help="the filled timed-session sheet")
    x.set_defaults(f=cmd_review_pack)
    x = s.add_parser("score")
    x.add_argument("--review")
    x.add_argument("--session")
    x.set_defaults(f=cmd_score)
    x = s.add_parser("viewer")
    x.add_argument("--out", required=True)
    x.add_argument("--banner", default="")
    x.set_defaults(f=cmd_viewer)
    a = p.parse_args(argv)
    a.f(a)


if __name__ == "__main__":
    main()
