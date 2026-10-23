#!/usr/bin/env python3
"""Extract the paper's gene lists from SI Dataset S1.

Usage: paper_gene_lists.py <Dataset_S1.xlsx> <out_dir>

Writes:
  paper_named_lists.tsv   the lists the paper reports (sheets "annotated DE all groups young/old"):
                          contrast, direction, list_header, list_n_in_header, gene_id, gene_name
  paper_deseq_all.parquet every row of the six DESeq2 sheets, with contrast
  paper_list_summary.json per contrast and direction: genes in the named list, the n written in
                          its header, and genes passing nominal P <= 0.05 and |log2FC| > 1 in the
                          DESeq2 sheet (all, and those with a gene name that is not the bare g ID)
"""
import json
import re
import sys
from pathlib import Path

import datetime

import openpyxl
import polars as pl


def text(v):
    """Cell as text. Excel turns some gene names (SEPT7, MARCH1, ...) into dates; those come back
    as datetime objects and are written as ISO dates, flagged by `excel_date_name` below."""
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()[:10]
    return str(v)

xlsx, out = sys.argv[1], Path(sys.argv[2])
out.mkdir(parents=True, exist_ok=True)
wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

contrast_of = {"1": "young_2h", "2": "young_24h", "3": "old_24h", "4": "young_13_49d",
               "5": "young_132d", "6": "old_41_93d"}
deseq_sheets = {"Group1_2Hr_young_DEseqres": "young_2h", "group2_24Hr_Young_DEseqresult": "young_24h",
                "group3_24Hr_OLD_DEseq": "old_24h", "Group4_13and49Days_Young_DEseq": "young_13_49d",
                "group5_132Days_young_DEseq": "young_132d", "group6_41to93days_oldDEseq": "old_41_93d"}

named = []
for sheet in ("annotated DE all groups young", "annotated DE all groups old"):
    rows = list(wb[sheet].iter_rows(values_only=True))
    header = rows[0]
    # columns come in (gi, "<group n- ... DE up|down n=NN>") pairs
    for j in range(0, len(header) - 1):
        h = header[j + 1]
        if header[j] != "gi" or not isinstance(h, str):
            continue
        m = re.search(r"group\s*(\d)\D.*DE\s+(up|down)\s+n\s*=\s*(\d+)", h, re.I)
        if not m:
            raise SystemExit(f"cannot parse list header {h!r} in {sheet}")
        g, direction, n = m.groups()
        for r in rows[1:]:
            gid = r[j]
            if gid is None or str(gid).strip() == "":
                continue
            named.append(dict(contrast=contrast_of[g], direction=direction.lower(), list_header=h.strip(),
                              list_n_in_header=int(n), gene_id=str(gid).strip(),
                              gene_name="" if r[j + 1] is None else text(r[j + 1]).strip()))
named_df = pl.DataFrame(named)
named_df.write_csv(out / "paper_named_lists.tsv", separator="\t")

frames = []
for sheet, contrast in deseq_sheets.items():
    rows = list(wb[sheet].iter_rows(values_only=True))
    cols = [c for c in rows[0] if c is not None]
    data = [dict(zip(cols, map(text, r[: len(cols)]))) for r in rows[1:] if r[0] is not None]
    df = pl.DataFrame(data, schema={c: pl.Utf8 for c in cols}).with_columns(pl.lit(contrast).alias("contrast"))
    df = df.with_columns([pl.col(c).cast(pl.Float64, strict=False) for c in
                          ("ControlMean", "TreatmentMean", "baseMean", "log2FoldChange", "lfcSE", "stat",
                           "pvalue", "padj")])
    frames.append(df.with_columns(pl.col("Gene_ID").cast(pl.Utf8), pl.col("Gene_Name").cast(pl.Utf8),
                                  pl.col("Gene_Note").cast(pl.Utf8)))
deseq = pl.concat(frames, how="diagonal_relaxed").with_columns(
    pl.col("Gene_Name").str.contains(r"^\d{4}-\d{2}-\d{2}$").fill_null(False).alias("excel_date_name"))
deseq.write_parquet(out / "paper_deseq_all.parquet")

summary = {"excel_date_gene_names": deseq.filter(pl.col("excel_date_name"))
           .select("Gene_ID", "Gene_Name", "Gene_Note").unique().sort("Gene_ID").rows()}
for contrast in contrast_of.values():
    d = deseq.filter(pl.col("contrast") == contrast)
    nominal = d.filter((pl.col("pvalue") <= 0.05) & (pl.col("log2FoldChange").abs() > 1))
    named_c = named_df.filter(pl.col("contrast") == contrast) if named_df.height else named_df
    for direction, sign in (("up", 1), ("down", -1)):
        nom = nominal.filter(pl.col("log2FoldChange") * sign > 1)
        nom_named = nom.filter(pl.col("Gene_Name") != pl.col("Gene_ID"))
        lst = named_c.filter(pl.col("direction") == direction) if named_c.height else named_c
        in_nominal = set(lst["gene_id"]) & set(nom["Gene_ID"]) if lst.height else set()
        summary[f"{contrast}_{direction}"] = dict(
            named_list_genes=lst.height,
            named_list_n_in_header=int(lst["list_n_in_header"][0]) if lst.height else None,
            deseq_nominal_p005_lfc1=nom.height,
            deseq_nominal_with_gene_name=nom_named.height,
            named_list_genes_found_in_nominal=len(in_nominal),
            deseq_padj005_lfc1=d.filter((pl.col("padj") <= 0.05) & (pl.col("log2FoldChange") * sign > 1)).height,
        )
(out / "paper_list_summary.json").write_text(json.dumps(summary, indent=1))
print(json.dumps(summary, indent=1))
